"""
Le a planilha de dados cadastrais dos clientes (Excel, uma aba por
empresa, colunas Empresa / Cliente / CNPJ / CEP / Endereco / Bairro /
Cidade / UF) e casa cada linha com um cliente ja cadastrado, para
completar o endereco. Usado pela tela de Clientes: primeiro mostra a
conferencia, depois aplica so as linhas que o escritorio aprovou.

A planilha tem dados de clientes (CNPJ/CPF) - ela so passa pela memoria
durante a importacao, nunca e gravada em disco.
"""

import re
import unicodedata
from difflib import SequenceMatcher
from io import BytesIO

from openpyxl import load_workbook

# nome normalizado do cabecalho -> campo
COLUNAS = {
    "EMPRESA": "empresa",
    "CLIENTE": "nome",
    "CNPJ": "cnpj",
    "CNPJ CPF": "cnpj",
    "CEP": "cep",
    "ENDERECO": "endereco",
    "BAIRRO": "bairro",
    "CIDADE": "cidade",
    "MUNICIPIO": "cidade",
    "UF": "uf",
    "ESTADO": "uf",
}
CAMPOS_ENDERECO = ["endereco", "bairro", "cidade", "uf", "cep"]
SEMELHANCA_MINIMA = 0.6


def normalizar(texto) -> str:
    """Maiusculo, sem acento e sem pontuacao: 'Câmara Mun. de Antônio' -> 'CAMARA MUN DE ANTONIO'."""
    texto = unicodedata.normalize("NFD", str(texto or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Z0-9]+", " ", texto.upper()).strip()


def _so_digitos(texto) -> str:
    return re.sub(r"\D", "", str(texto or ""))


def _limpar(valor) -> str | None:
    if valor is None:
        return None
    texto = re.sub(r"\s+", " ", str(valor).replace("\xa0", " ")).strip()
    return texto or None


def _formatar_cep(valor) -> str | None:
    digitos = _so_digitos(valor)
    if isinstance(valor, (int, float)) and len(digitos) == 7:
        digitos = "0" + digitos  # Excel come o zero da frente
    return f"{digitos[:5]}-{digitos[5:]}" if len(digitos) == 8 else None


def ler_planilha(conteudo: bytes) -> list[dict]:
    """Devolve uma linha por cliente da planilha, de todas as abas."""
    livro = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    linhas = []
    for aba in livro.worksheets:
        mapa_colunas = None
        for numero, valores in enumerate(aba.iter_rows(values_only=True), start=1):
            if mapa_colunas is None:
                cabecalho = {i: COLUNAS.get(normalizar(v)) for i, v in enumerate(valores)}
                if "nome" in cabecalho.values():
                    mapa_colunas = {i: campo for i, campo in cabecalho.items() if campo}
                continue
            dados = {campo: valores[i] if i < len(valores) else None for i, campo in mapa_colunas.items()}
            nome = _limpar(dados.get("nome"))
            if not nome:
                continue
            linhas.append(
                {
                    "linha_id": f"{aba.title}!{numero}",
                    "aba": aba.title,
                    "numero": numero,
                    "empresa": _limpar(dados.get("empresa")) or aba.title,
                    "nome": nome,
                    "cnpj": _limpar(dados.get("cnpj")),
                    "cep": _formatar_cep(dados.get("cep")),
                    "endereco": _limpar(dados.get("endereco")),
                    "bairro": _limpar(dados.get("bairro")),
                    "cidade": _limpar(dados.get("cidade")),
                    "uf": (_limpar(dados.get("uf")) or "").upper()[:2] or None,
                }
            )
    return linhas


def _semelhanca(a: str, b: str) -> float:
    if a in b or b in a:
        # 'SELMA' x 'SELMA ED GRASIELA': um contido no outro ja e um bom sinal
        return max(0.85, SequenceMatcher(None, a, b).ratio())
    return SequenceMatcher(None, a, b).ratio()


def casar_com_clientes(linhas: list[dict], empresas: list, clientes: list) -> tuple[list[dict], list[dict]]:
    """
    Casa cada linha com um cliente da mesma empresa: primeiro pelo nome
    igual, depois pelo CNPJ (se so um cliente da empresa tiver aquele
    CNPJ), depois pelo nome igual em outra empresa do grupo e, por
    ultimo, pelo nome parecido. Um cliente nunca e casado com duas
    linhas. Devolve (casadas, nao_encontradas).
    """
    empresa_por_nome = {normalizar(e.nome): e for e in empresas}
    clientes_por_empresa: dict[int, list] = {}
    for c in clientes:
        clientes_por_empresa.setdefault(c.empresa_id, []).append(c)

    casadas, nao_encontradas, pendentes = [], [], []
    usados: set[int] = set()

    for linha in linhas:
        empresa = empresa_por_nome.get(normalizar(linha["empresa"])) or empresa_por_nome.get(normalizar(linha["aba"]))
        if empresa is None:
            nao_encontradas.append({**linha, "motivo": f"Empresa \"{linha['empresa']}\" não cadastrada"})
            continue
        candidatos = clientes_por_empresa.get(empresa.id, [])
        nome = normalizar(linha["nome"])

        iguais = [c for c in candidatos if normalizar(c.nome) == nome and c.id not in usados]
        if iguais:
            usados.add(iguais[0].id)
            casadas.append({**linha, "cliente": iguais[0], "metodo": "nome", "semelhanca": 1.0})
            continue

        digitos = _so_digitos(linha["cnpj"])
        mesmo_cnpj = [c for c in candidatos if digitos and _so_digitos(c.cnpj) == digitos]
        if len(mesmo_cnpj) == 1 and mesmo_cnpj[0].id not in usados:
            usados.add(mesmo_cnpj[0].id)
            casadas.append({**linha, "cliente": mesmo_cnpj[0], "metodo": "cnpj", "semelhanca": 1.0})
            continue

        # cadastrado em outra empresa do grupo: so vale nome igual e unico
        outra_empresa = [c for c in clientes if normalizar(c.nome) == nome and c.id not in usados]
        if len(outra_empresa) == 1:
            usados.add(outra_empresa[0].id)
            casadas.append({**linha, "cliente": outra_empresa[0], "metodo": "outra_empresa", "semelhanca": 1.0})
            continue

        pendentes.append((linha, candidatos, nome))

    # nome parecido: as melhores semelhancas escolhem primeiro
    sugestoes = []
    for linha, candidatos, nome in pendentes:
        notas = sorted(
            ((_semelhanca(nome, normalizar(c.nome)), c) for c in candidatos if c.id not in usados),
            key=lambda par: par[0],
            reverse=True,
        )
        sugestoes.append((notas[0] if notas else (0.0, None), linha))
    sugestoes.sort(key=lambda s: s[0][0], reverse=True)
    for (nota, cliente), linha in sugestoes:
        if cliente is not None and nota >= SEMELHANCA_MINIMA and cliente.id not in usados:
            usados.add(cliente.id)
            casadas.append({**linha, "cliente": cliente, "metodo": "parecido", "semelhanca": round(nota, 2)})
        else:
            nao_encontradas.append({**linha, "motivo": "Nenhum cliente cadastrado com nome parecido"})

    casadas.sort(key=lambda c: (c["aba"], c["numero"]))
    nao_encontradas.sort(key=lambda c: (c["aba"], c["numero"]))
    return casadas, nao_encontradas


def mudancas_da_linha(linha: dict) -> dict:
    """Campos de endereco que a planilha preenche ou corrige: {campo: [atual, novo]}."""
    cliente = linha["cliente"]
    mudancas = {}
    for campo in CAMPOS_ENDERECO:
        novo = linha[campo]
        atual = getattr(cliente, campo)
        if novo and (atual or "").strip() != novo:
            mudancas[campo] = [atual, novo]
    if linha["cnpj"] and not cliente.cnpj:
        mudancas["cnpj"] = [None, linha["cnpj"]]
    return mudancas
