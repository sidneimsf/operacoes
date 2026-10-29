"""
Movimentacao de cobranca mensal (menu Gerencia): o que antes era a planilha
"MOVIMENTACAO <MES> <ANO>.xlsx" - uma aba por empresa, uma linha por cliente,
com ISS e INSS retidos calculados pelas aliquotas do mes.

Aqui ficam as regras de calculo (usadas pela tela, pelo Excel e pelo PDF),
a geracao do Excel no mesmo layout da planilha antiga (com formulas, pra
quem abrir no Excel continuar vendo as contas) e a leitura da planilha
antiga pra importar um mes ja preenchido.
"""

import re
import unicodedata
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MESES = [
    "JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
    "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO",
]

# tipo de servico -> codigos da nota. A descricao recebe o mes/ano sozinha.
# Por padrao limpeza (cessao de mao de obra) tem ISS e INSS retidos; apoio nao.
TIPOS_SERVICO = [
    {
        "chave": "limpeza",
        "label": "Limpeza",
        "ctn": "07.10.02",
        "nbs": "118031000",
        "descricao": "SERVIÇOS DE LIMPEZA EM {mes}/{ano}",
        "retem_iss": True,
        "retem_inss": True,
    },
    {
        "chave": "apoio",
        "label": "Apoio (combinados)",
        "ctn": "17.02.02",
        "nbs": "118064000",
        "descricao": "SERVIÇOS COMBINADOS DE APOIO EM {mes}/{ano}",
        "retem_iss": False,
        "retem_inss": False,
    },
]
TIPOS_SERVICO_POR_CHAVE = {t["chave"]: t for t in TIPOS_SERVICO}
INFORMACOES_COMPLEMENTARES = "EMPRESA OPTANTE PELO SIMPLES NACIONAL."
ALIQUOTA_INSS_PADRAO = 11.0
MAXIMO_REMESSA = 6


def nome_mes(mes: int) -> str:
    return MESES[mes - 1]


def descricao_servico(tipo_servico: str, ano: int, mes: int) -> str:
    tipo = TIPOS_SERVICO_POR_CHAVE.get(tipo_servico, TIPOS_SERVICO[0])
    return tipo["descricao"].format(mes=nome_mes(mes), ano=ano)


def valor_nota(composicao: list | None) -> float:
    return round(sum(float(item.get("valor") or 0) for item in (composicao or [])), 2)


def calcular_linha(linha, cabecalho) -> dict:
    """Valores calculados de uma linha com as aliquotas do mes (cabecalho)."""
    nota = valor_nota(linha.composicao)
    iss = round(nota * (cabecalho.aliquota_iss or 0) / 100, 2) if linha.retem_iss else 0.0
    inss = round(nota * (cabecalho.aliquota_inss or 0) / 100, 2) if linha.retem_inss else 0.0
    tipo = linha.tipo_servico
    simples = cabecalho.simples_limpeza if tipo == "limpeza" else cabecalho.simples_apoio
    return {
        "valor_nota": nota,
        "valor_iss": iss,
        "valor_inss": inss,
        "valor_boleto": round(nota - iss - inss, 2),
        "aliquota_simples": simples,
    }


def formatar_moeda(valor: float) -> str:
    texto = f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {texto}"


def texto_composicao(composicao: list | None) -> str:
    """'Valor mensal R$ 1.693,51 + Desconto falta - R$ 188,17' quando ha mais de um item."""
    itens = composicao or []
    if len(itens) < 2:
        return ""
    partes = []
    for item in itens:
        valor = float(item.get("valor") or 0)
        sinal = "- " if valor < 0 else ""
        descricao = (item.get("descricao") or "").strip()
        partes.append(f"{descricao} {sinal}{formatar_moeda(abs(valor))}".strip())
    return " + ".join(partes)


def observacao_completa(linha) -> str:
    partes = [linha.observacao, linha.observacao_mes]
    return " - ".join(p.strip() for p in partes if p and p.strip())


# ---------------------------------------------------------------------------
# Excel no layout da planilha antiga
# ---------------------------------------------------------------------------

CABECALHO_COLUNAS = [
    "REMESSA", "CLIENTE", "CNPJ", "MUNICÍPIO", "CTN", "DESCRIÇÃO DO SERVIÇO", "NBS",
    "INFORMAÇÕES COMPLEMENTARES", "ISS", "VALOR RETIDO CP (INSS)", "ALÍQUOTA SIMPLES %",
    "VALOR NOTA", "VALOR BOLETO", "Nº DA NOTA", "VENCIMENTO", "BOLETO", "CLIENTE",
    "DOCUMENTOS", "OBSERVAÇÃO",
]
LARGURAS_COLUNAS = [8, 42, 21, 18, 10, 44, 12, 38, 14, 16, 12, 16, 16, 9, 8, 8, 8, 11, 90]
FORMATO_MOEDA = '"R$"\\ #,##0.00'


def _numero_formula(valor: float) -> str:
    return f"{valor:.2f}".rstrip("0").rstrip(".")


def _formula_valor_nota(composicao: list | None):
    itens = [float(i.get("valor") or 0) for i in (composicao or [])]
    if len(itens) <= 1:
        return round(itens[0], 2) if itens else 0
    formula = "=" + _numero_formula(itens[0])
    for v in itens[1:]:
        formula += ("-" if v < 0 else "+") + _numero_formula(abs(v))
    return formula


def _percentual_formula(valor: float) -> str:
    return _numero_formula(valor or 0) + "%"


def gerar_excel(cabecalhos: list) -> bytes:
    """Uma aba por empresa (cabecalhos ja carregados com empresa e linhas)."""
    wb = Workbook()
    wb.remove(wb.active)

    negrito = Font(name="Calibri", size=10, bold=True)
    normal = Font(name="Calibri", size=10)
    fundo_cabecalho = PatternFill("solid", fgColor="D9E1F2")
    fino = Side(style="thin", color="999999")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cab in cabecalhos:
        ws = wb.create_sheet(title=cab.empresa.nome[:31])
        mes_texto = nome_mes(cab.mes)

        ws.merge_cells("A1:S1")
        ws.merge_cells("A2:S2")
        ws["A1"] = f"COBRANÇA SISTEMA SICREDI - {cab.empresa.nome}"
        ws["A2"] = f"MOVIMENTAÇÃO DE COBRANÇA {mes_texto} - REF. SERVIÇOS PRESTADOS EM {mes_texto}/{cab.ano}"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True)
        ws["A2"].font = Font(name="Calibri", size=12, bold=True)
        ws["A1"].alignment = ws["A2"].alignment = Alignment(horizontal="center")

        titulos = list(CABECALHO_COLUNAS)
        titulos[8] = f"ISS {_numero_formula(cab.aliquota_iss or 0).replace('.', ',')}%"
        for col, titulo in enumerate(titulos, start=1):
            celula = ws.cell(row=3, column=col, value=titulo)
            celula.font = negrito
            celula.fill = fundo_cabecalho
            celula.alignment = centro
            celula.border = borda

        linhas = sorted(cab.linhas, key=lambda l: (l.remessa, l.ordem, l.id))
        primeira = 4
        for idx, linha in enumerate(linhas):
            r = primeira + idx
            calc = calcular_linha(linha, cab)
            tipo = TIPOS_SERVICO_POR_CHAVE.get(linha.tipo_servico, TIPOS_SERVICO[0])
            celula_iss = f"=L{r}*{_percentual_formula(cab.aliquota_iss)}" if linha.retem_iss else "X"
            celula_inss = f"=L{r}*{_percentual_formula(cab.aliquota_inss)}" if linha.retem_inss else "X"
            descontos = [f"I{r}"] * linha.retem_iss + [f"J{r}"] * linha.retem_inss
            celula_boleto = f"=L{r}" + "".join(f"-{c}" for c in descontos)
            documentos = ("OK" if linha.documentos_ok else "X") if linha.exige_documentos else None
            valores = [
                f"{linha.remessa}º",
                linha.nome_cliente,
                linha.cnpj,
                linha.municipio,
                tipo["ctn"],
                descricao_servico(linha.tipo_servico, cab.ano, cab.mes),
                int(tipo["nbs"]),
                INFORMACOES_COMPLEMENTARES,
                celula_iss,
                celula_inss,
                calc["aliquota_simples"],
                _formula_valor_nota(linha.composicao),
                celula_boleto,
                linha.numero_nota,
                linha.vencimento,
                "OK" if linha.boleto_ok else "X",
                "OK" if linha.cliente_ok else "X",
                documentos,
                observacao_completa(linha) or None,
            ]
            for col, valor in enumerate(valores, start=1):
                celula = ws.cell(row=r, column=col, value=valor)
                celula.font = normal
                celula.border = borda
                if col in (9, 10, 12, 13):
                    celula.number_format = FORMATO_MOEDA
                if col in (1, 5, 7, 9, 10, 11, 14, 15, 16, 17, 18):
                    celula.alignment = Alignment(horizontal="center")

        total = primeira + len(linhas)
        if linhas:
            for col in (9, 10, 12, 13):
                letra = get_column_letter(col)
                celula = ws.cell(row=total, column=col, value=f"=SUM({letra}{primeira}:{letra}{total - 1})")
                celula.font = negrito
                celula.number_format = FORMATO_MOEDA
                celula.border = borda
            ws.cell(row=total, column=2, value="TOTAL").font = negrito

        for col, largura in enumerate(LARGURAS_COLUNAS, start=1):
            ws.column_dimensions[get_column_letter(col)].width = largura
        ws.row_dimensions[3].height = 30
        ws.freeze_panes = "C4"

    if not wb.worksheets:
        wb.create_sheet(title="Movimentação")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Importacao da planilha antiga
# ---------------------------------------------------------------------------

def normalizar(texto) -> str:
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texto).strip().upper()


def _limpar(valor) -> str | None:
    if valor is None:
        return None
    texto = re.sub(r"\s+", " ", str(valor)).strip()
    return texto or None


def _percentual_da_formula(formula) -> float | None:
    """'=L4*4.62%' -> 4.62"""
    achado = re.search(r"\*\s*([\d.,]+)\s*%", str(formula or ""))
    return float(achado.group(1).replace(",", ".")) if achado else None


def _composicao_da_celula(formula, valor_calculado) -> list[dict]:
    """'=5898.28+6048.29-188.17' vira um item por parcela; valores negativos
    (descontos) ficam como do mes, os positivos como fixos."""
    if isinstance(formula, (int, float)):
        return [{"descricao": "Valor mensal", "valor": round(float(formula), 2), "fixo": True}]
    texto = str(formula or "").strip()
    if texto.startswith("=") and re.fullmatch(r"=\s*[\d.\s+-]+", texto):
        parcelas = re.findall(r"([+-]?)\s*([\d.]+)", texto[1:])
        itens = []
        for sinal, numero in parcelas:
            valor = float(numero) * (-1 if sinal == "-" else 1)
            if not itens:
                itens.append({"descricao": "Valor mensal", "valor": round(valor, 2), "fixo": True})
            elif valor < 0:
                itens.append({"descricao": "Desconto", "valor": round(valor, 2), "fixo": False})
            else:
                itens.append({"descricao": "Adicional", "valor": round(valor, 2), "fixo": True})
        if itens:
            return itens
    if isinstance(valor_calculado, (int, float)):
        return [{"descricao": "Valor mensal", "valor": round(float(valor_calculado), 2), "fixo": True}]
    return []


def _remessa(valor) -> int:
    achado = re.search(r"\d+", str(valor or ""))
    numero = int(achado.group()) if achado else 1
    return min(max(numero, 1), MAXIMO_REMESSA)


def _vencimento(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return str(int(valor)).zfill(2)
    return _limpar(valor)


def _ok(valor) -> bool:
    return normalizar(valor) == "OK"


def _tipo_servico(ctn, descricao) -> str:
    ctn = str(ctn or "").strip()
    for tipo in TIPOS_SERVICO:
        if tipo["ctn"] == ctn:
            return tipo["chave"]
    return "limpeza" if "LIMPEZA" in normalizar(descricao) else "apoio"


def ler_planilha(conteudo: bytes) -> list[dict]:
    """Uma entrada por aba: nome da aba, aliquotas encontradas e linhas."""
    try:
        livro = load_workbook(BytesIO(conteudo))
        livro_valores = load_workbook(BytesIO(conteudo), data_only=True)
    except Exception:
        raise ValueError("Não foi possível ler o arquivo. Envie a planilha no formato .xlsx.")

    abas = []
    for ws in livro.worksheets:
        ws_valores = livro_valores[ws.title]
        aliquota_iss = None
        aliquota_inss = None
        simples = {}
        linhas = []

        # o cabecalho "ISS 4,62%" ja traz a aliquota
        cabecalho_iss = ws.cell(row=3, column=9).value
        achado = re.search(r"([\d]+[.,]?\d*)\s*%", str(cabecalho_iss or ""))
        if achado:
            aliquota_iss = float(achado.group(1).replace(",", "."))

        for r in range(4, ws.max_row + 1):
            nome = _limpar(ws.cell(row=r, column=2).value)
            celula_nota = ws.cell(row=r, column=12).value
            if not nome or celula_nota is None or normalizar(nome) == "TOTAL":
                continue
            if str(celula_nota).upper().startswith("=SUM"):
                continue

            formula_iss = ws.cell(row=r, column=9).value
            formula_inss = ws.cell(row=r, column=10).value
            retem_iss = str(formula_iss or "").startswith("=")
            retem_inss = str(formula_inss or "").startswith("=")
            aliquota_iss = aliquota_iss or _percentual_da_formula(formula_iss)
            aliquota_inss = aliquota_inss or _percentual_da_formula(formula_inss)

            tipo = _tipo_servico(ws.cell(row=r, column=5).value, ws.cell(row=r, column=6).value)
            aliquota_simples = ws.cell(row=r, column=11).value
            if isinstance(aliquota_simples, (int, float)):
                simples.setdefault(tipo, float(aliquota_simples))

            numero_nota = ws.cell(row=r, column=14).value
            try:
                numero_nota = int(numero_nota) if numero_nota not in (None, "") else None
            except (TypeError, ValueError):
                numero_nota = None

            documentos = normalizar(ws.cell(row=r, column=18).value)
            cnpj = ws.cell(row=r, column=3).value
            linhas.append(
                {
                    "remessa": _remessa(ws.cell(row=r, column=1).value),
                    "nome_cliente": nome[:200],
                    "cnpj": (_limpar(cnpj) or "")[:20] or None,
                    "municipio": (_limpar(ws.cell(row=r, column=4).value) or "")[:100] or None,
                    "tipo_servico": tipo,
                    "retem_iss": retem_iss,
                    "retem_inss": retem_inss,
                    "composicao": _composicao_da_celula(celula_nota, ws_valores.cell(row=r, column=12).value),
                    "numero_nota": numero_nota,
                    "vencimento": (_vencimento(ws.cell(row=r, column=15).value) or "")[:20] or None,
                    "boleto_ok": _ok(ws.cell(row=r, column=16).value),
                    "cliente_ok": _ok(ws.cell(row=r, column=17).value),
                    "exige_documentos": bool(documentos),
                    "documentos_ok": documentos == "OK",
                    "observacao": (_limpar(ws.cell(row=r, column=19).value) or "")[:1000] or None,
                }
            )

        if linhas:
            abas.append(
                {
                    "aba": ws.title,
                    "aliquota_iss": aliquota_iss or 0.0,
                    "aliquota_inss": aliquota_inss or ALIQUOTA_INSS_PADRAO,
                    "simples_limpeza": simples.get("limpeza"),
                    "simples_apoio": simples.get("apoio"),
                    "linhas": linhas,
                }
            )
    return abas
